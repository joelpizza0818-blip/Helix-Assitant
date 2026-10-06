import asyncio
import logging
import csv
import json
import os
from typing import Any, Callable, Dict, List, Union

logger = logging.getLogger(__name__)

class DocumentTool:
    """Tool for reading and writing documents."""
    
    async def read_pdf(self, path: str) -> str:
        """Reads text from a PDF file using PyMuPDF (fitz)."""
        logger.info(f"Reading PDF: {path}")
        def _read():
            import fitz
            doc = fitz.open(path)
            text = ""
            for page in doc:
                text += page.get_text()
            doc.close()
            return text
        return await asyncio.to_thread(_read)

    async def read_docx(self, path: str) -> str:
        """Reads text from a DOCX file using python-docx."""
        logger.info(f"Reading DOCX: {path}")
        def _read():
            import docx
            doc = docx.Document(path)
            return "\n".join([p.text for p in doc.paragraphs])
        return await asyncio.to_thread(_read)

    async def read_txt(self, path: str) -> str:
        """Reads text from a TXT file."""
        logger.info(f"Reading TXT: {path}")
        def _read():
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        return await asyncio.to_thread(_read)

    async def read_csv(self, path: str) -> List[Dict[str, Any]]:
        """Reads data from a CSV file."""
        logger.info(f"Reading CSV: {path}")
        def _read():
            results = []
            with open(path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    results.append(dict(row))
            return results
        return await asyncio.to_thread(_read)

    async def read_json(self, path: str) -> Union[Dict, List]:
        """Reads data from a JSON file."""
        logger.info(f"Reading JSON: {path}")
        def _read():
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        return await asyncio.to_thread(_read)

    async def write_docx(self, path: str, content: str) -> None:
        """Writes text to a DOCX file."""
        logger.info(f"Writing to DOCX: {path}")
        def _write():
            import docx
            doc = docx.Document()
            doc.add_paragraph(content)
            doc.save(path)
        await asyncio.to_thread(_write)

    async def summarize_document(self, path: str, model_call_fn: Callable) -> str:
        """Summarizes a document using a model."""
        logger.info(f"Summarizing document: {path}")
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            content = await self.read_pdf(path)
        elif ext == ".docx":
            content = await self.read_docx(path)
        else:
            content = await self.read_txt(path)
            
        prompt = f"Please summarize the following document:\n\n{content}"
        return await model_call_fn(prompt)

    async def answer_question_in_document(self, path: str, question: str, model_call_fn: Callable) -> str:
        """Answers a question based on a document using a model."""
        logger.info(f"Answering question about document {path}: {question}")
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            content = await self.read_pdf(path)
        elif ext == ".docx":
            content = await self.read_docx(path)
        else:
            content = await self.read_txt(path)
            
        prompt = f"Based on the following document, answer this question: {question}\n\nDocument:\n{content}"
        return await model_call_fn(prompt)
